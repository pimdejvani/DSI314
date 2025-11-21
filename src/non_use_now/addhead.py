from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

# ===== ปรับค่าตามโจทย์ของคุณ =====
path_doc = Path(r"C:\Users\Pimdej\DSI314\ตัวอย่างหลักสูตร (เคสยาก)\4. อว.รม. ความสัมพันธ์ระหว่างประเทศ (as of Aug 4).docx")
target_text = "หมวดที่ 2 คุณสมบัติผู้เข้าศึกษา"
output_dir = Path(r"C:\Users\Pimdej\DSI314\ตัวอย่างหลักสูตร (เคสยาก)")
insert_only_first_match = True  # True = ใส่ก่อนครั้งแรกที่เจอ, False = ใส่ก่อนทุกครั้งที่เจอ
# ==================================

def insert_page_break_before_paragraph(paragraph) -> None:
    """
    แทรก <w:br w:type="page"/> เป็น run ตัวแรกของย่อหน้า
    """
    p = paragraph._p
    r = OxmlElement('w:r')
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    r.append(br)
    # แทรกเป็นลูกลำดับแรกใน <w:p>
    p.insert(0, r)

def left_align_all_text(doc: Document) -> int:
    """
    จัดชิดซ้ายให้ทุกย่อหน้าที่เข้าถึงได้:
    - เนื้อหาหลัก (doc.paragraphs)
    - ย่อหน้าในตารางทั้งหมด
    - ย่อหน้าใน header/footer ของทุก section
    คืนค่าจำนวนย่อหน้าที่ปรับแล้ว
    """
    changed = 0

    # 1) ย่อหน้าทั่วไปในเนื้อหา
    for p in doc.paragraphs:
        if p.alignment != WD_ALIGN_PARAGRAPH.LEFT:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            changed += 1

    # 2) ย่อหน้าในตารางทั้งหมด
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if p.alignment != WD_ALIGN_PARAGRAPH.LEFT:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        changed += 1

    # 3) Header/Footer ทุก section
    for section in doc.sections:
        # header
        for p in section.header.paragraphs:
            if p.alignment != WD_ALIGN_PARAGRAPH.LEFT:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                changed += 1
        for table in section.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if p.alignment != WD_ALIGN_PARAGRAPH.LEFT:
                            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                            changed += 1
        # footer
        for p in section.footer.paragraphs:
            if p.alignment != WD_ALIGN_PARAGRAPH.LEFT:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                changed += 1
        for table in section.footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if p.alignment != WD_ALIGN_PARAGRAPH.LEFT:
                            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                            changed += 1

    return changed

def main():
    if not path_doc.exists():
        raise FileNotFoundError(f"ไม่พบไฟล์: {path_doc}")

    if path_doc.suffix.lower() != ".docx":
        raise ValueError("สคริปต์นี้รองรับเฉพาะ .docx (ถ้าเป็น .doc กรุณาแปลงเป็น .docx ก่อน)")

    doc = Document(str(path_doc))

    # 1) แทรก page break ก่อนย่อหน้าที่มี target_text
    found = 0
    for para in doc.paragraphs:
        text = (para.text or "").strip()
        if target_text in text:
            insert_page_break_before_paragraph(para)
            found += 1
            if insert_only_first_match:
                break

    if found == 0:
        print(f"ไม่พบข้อความเป้าหมายในไฟล์: {target_text!r}")
    else:
        print(f"พบข้อความเป้าหมาย {found} ครั้ง และได้แทรก page break {'ครั้งแรกเท่านั้น' if insert_only_first_match else 'ก่อนทุกครั้งที่พบ'}")

    # 2) จัดชิดซ้ายทุกข้อความที่เข้าถึงได้
    changed = left_align_all_text(doc)
    print(f"จัดชิดซ้ายย่อหน้าทั้งหมด: {changed} ย่อหน้า")

    # ตั้งชื่อไฟล์ใหม่: เติม "(แก้) " หน้าชื่อเดิม
    new_name = "(แก้) " + path_doc.name
    out_path = (output_dir if output_dir else path_doc.parent) / new_name

    # บันทึก
    doc.save(str(out_path))
    print(f"บันทึกไฟล์ใหม่แล้วที่: {out_path}")

if __name__ == "__main__":
    main()
